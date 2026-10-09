import os
import io
import zlib
import zipfile
import base64
import xml.etree.ElementTree as ET
from typing import Tuple, List, Dict, Any, Optional

# PDF
try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

# DOCX
try:
    import docx
except ImportError:
    docx = None

# PPTX
try:
    from pptx import Presentation
except ImportError:
    Presentation = None

# HWP (OleFile)
try:
    import olefile
except ImportError:
    olefile = None

# PIL
try:
    from PIL import Image
except ImportError:
    Image = None


def parse_page_ranges(range_str: Optional[str], total_count: int) -> List[int]:
    """
    "1-3, 5, 8-10" 또는 "2" 형태의 문자열을 받아 0-indexed 정수 리스트를 반환합니다.
    빈 문자열이거나 유효하지 않으면 range(total_count) 전체를 반환합니다.
    """
    if not range_str or not range_str.strip():
        return list(range(total_count))
    
    selected = set()
    parts = range_str.replace(" ", "").split(",")
    for part in parts:
        if not part:
            continue
        if "-" in part:
            bounds = part.split("-")
            start = int(bounds[0]) if bounds[0].isdigit() else 1
            end = int(bounds[1]) if len(bounds) > 1 and bounds[1].isdigit() else total_count
            for p in range(start, end + 1):
                if 1 <= p <= total_count:
                    selected.add(p - 1)
        elif part.isdigit():
            p = int(part)
            if 1 <= p <= total_count:
                selected.add(p - 1)
                
    if not selected:
        return list(range(total_count))
    return sorted(list(selected))


def extract_content_from_file(
    file_path: str, 
    filename: str, 
    page_range: Optional[str] = None
) -> Dict[str, Any]:
    """
    파일 확장자에 맞춰 텍스트 및 이미지 데이터를 추출합니다.
    지정된 페이지/슬라이드 범위(page_range)가 있을 경우 해당 범위만 추출합니다.
    """
    ext = os.path.splitext(filename)[1].lower()
    
    result = {
        "text": "",
        "images": [],
        "filename": filename,
        "format": ext.replace(".", ""),
        "page_info": ""
    }

    if ext == ".pdf":
        result["text"], result["images"], result["page_info"] = parse_pdf(file_path, page_range)
    elif ext == ".docx":
        result["text"], result["page_info"] = parse_docx(file_path, page_range)
    elif ext == ".pptx":
        result["text"], result["page_info"] = parse_pptx(file_path, page_range)
    elif ext == ".hwpx":
        result["text"], result["page_info"] = parse_hwpx(file_path, page_range)
    elif ext == ".hwp":
        result["text"], result["page_info"] = parse_hwp(file_path, page_range)
    elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
        result["images"] = parse_image(file_path, ext)
        result["text"] = f"[이미지 파일: {filename}]"
        result["page_info"] = "단일 이미지 (1페이지)"
    elif ext in [".txt", ".md", ".csv", ".json"]:
        result["text"] = parse_plain_text(file_path)
        result["page_info"] = "텍스트 문서"
    else:
        try:
            result["text"] = parse_plain_text(file_path)
            result["page_info"] = "일반 문서"
        except Exception:
            raise ValueError(f"지원되지 않는 파일 형식입니다: {ext}")

    return result


def parse_pdf(file_path: str, page_range: Optional[str] = None) -> Tuple[str, List[Dict[str, str]], str]:
    if not fitz:
        raise ImportError("PyMuPDF (fitz) 라이브러리가 설치되어 있지 않습니다.")
    
    doc = fitz.open(file_path)
    total_pages = len(doc)
    selected_indices = parse_page_ranges(page_range, total_pages)
    
    text_chunks = []
    images = []
    
    for idx in selected_indices:
        page = doc[idx]
        page_text = page.get_text("text")
        if page_text.strip():
            text_chunks.append(f"--- [PDF 페이지 {idx + 1} / 전체 {total_pages}] ---\n" + page_text)
        
        # 텍스트가 적거나 스캔된 PDF인 경우 이미지 렌더링
        if len(page_text.strip()) < 50 and len(images) < 3:
            pix = page.get_pixmap(dpi=150)
            img_bytes = pix.tobytes("jpeg")
            b64_str = base64.b64encode(img_bytes).decode("utf-8")
            images.append({
                "mime_type": "image/jpeg",
                "data": b64_str
            })

    doc.close()
    page_info = f"전체 {total_pages}페이지 중 {len(selected_indices)}개 페이지 추출됨 ({', '.join(str(i+1) for i in selected_indices)}페이지)"
    return "\n\n".join(text_chunks), images, page_info


def parse_pptx(file_path: str, page_range: Optional[str] = None) -> Tuple[str, str]:
    if not Presentation:
        raise ImportError("python-pptx 라이브러리가 설치되어 있지 않습니다.")
    
    prs = Presentation(file_path)
    total_slides = len(prs.slides)
    selected_indices = parse_page_ranges(page_range, total_slides)
    
    chunks = []
    for idx in selected_indices:
        slide = prs.slides[idx]
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        slide_text.append(" | ".join(row_text))
        
        if slide_text:
            chunks.append(f"--- [슬라이드 {idx + 1} / 전체 {total_slides}] ---\n" + "\n".join(slide_text))
            
    page_info = f"전체 {total_slides}개 슬라이드 중 {len(selected_indices)}개 슬라이드 추출됨"
    return "\n\n".join(chunks), page_info


def parse_docx(file_path: str, page_range: Optional[str] = None) -> Tuple[str, str]:
    if not docx:
        raise ImportError("python-docx 라이브러리가 설치되어 있지 않습니다.")
    
    doc = docx.Document(file_path)
    # DOCX는 섹션/문단 단위로 추출
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    # 표 내용도 추출
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs.append(" | ".join(row_text))
                
    total_count = len(paragraphs)
    if total_count == 0:
        return "", "문서가 비어 있습니다."
        
    # 만약 페이지 범위가 주어졌다면 문단을 분할해 선택
    if page_range and page_range.strip():
        # 대략 5문단 = 1페이지로 추정 분할
        est_pages = max(1, (total_count + 4) // 5)
        selected_pages = parse_page_ranges(page_range, est_pages)
        filtered = []
        for p_idx in selected_pages:
            start_p = p_idx * 5
            end_p = min(total_count, (p_idx + 1) * 5)
            filtered.extend(paragraphs[start_p:end_p])
        page_info = f"Word 문서 약 {est_pages}페이지 분량 중 {len(selected_pages)}페이지 분량 추출"
        return "\n\n".join(filtered), page_info
    
    return "\n\n".join(paragraphs), f"Word 전체 내용 ({total_count}개 단락) 추출"


def parse_hwpx(file_path: str, page_range: Optional[str] = None) -> Tuple[str, str]:
    chunks = []
    try:
        with zipfile.ZipFile(file_path, 'r') as zf:
            section_files = [f for f in zf.namelist() if f.startswith('Contents/section') and f.endswith('.xml')]
            section_files.sort()
            total_sections = len(section_files)
            
            selected_sections = parse_page_ranges(page_range, max(1, total_sections))
            
            for idx in selected_sections:
                if idx < total_sections:
                    s_file = section_files[idx]
                    xml_data = zf.read(s_file)
                    root = ET.fromstring(xml_data)
                    sec_texts = []
                    for elem in root.iter():
                        if elem.text and elem.text.strip():
                            sec_texts.append(elem.text.strip())
                    if sec_texts:
                        chunks.append(f"--- [HWPX 구역 {idx + 1}] ---\n" + "\n".join(sec_texts))
                        
        page_info = f"HWPX 문서 {total_sections}개 구역 중 {len(chunks)}개 구역 추출"
        return "\n\n".join(chunks), page_info
    except Exception as e:
        return f"[HWPX 파싱 오류: {str(e)}]", "오류 발생"


def parse_hwp(file_path: str, page_range: Optional[str] = None) -> Tuple[str, str]:
    if not olefile:
        raise ImportError("olefile 라이브러리가 설치되어 있지 않습니다.")
    
    try:
        ole = olefile.OleFileIO(file_path)
        sections = []
        for entry in ole.listdir():
            if len(entry) == 2 and entry[0] == 'BodyText' and entry[1].startswith('Section'):
                sections.append(entry)
                
        total_sections = len(sections)
        selected_indices = parse_page_ranges(page_range, max(1, total_sections)) if total_sections > 0 else [0]
        
        chunks = []
        for idx in selected_indices:
            if idx < total_sections:
                sec = sections[idx]
                stream = ole.openstream(sec).read()
                try:
                    decompressed = zlib.decompress(stream, -15)
                    text = decompressed.decode('utf-16-le', errors='ignore')
                except Exception:
                    try:
                        text = stream.decode('utf-16-le', errors='ignore')
                    except Exception:
                        text = ""
                
                filtered_text = "".join(c for c in text if c.isprintable() or c in "\n\r\t ")
                if filtered_text.strip():
                    chunks.append(f"--- [HWP 섹션 {idx + 1}] ---\n" + filtered_text.strip())
                    
        ole.close()
        if chunks:
            page_info = f"HWP 문서 {total_sections}개 섹션 중 {len(chunks)}개 섹션 추출"
            return "\n\n".join(chunks), page_info
            
        # PrvText fallback
        if ole.exists('PrvText'):
            data = ole.openstream('PrvText').read()
            text = data.decode('utf-16-le', errors='ignore')
            return text.strip(), "HWP 미리보기 텍스트 추출"
            
        return "[HWP 파일에서 추출 가능한 텍스트를 찾지 못했습니다.]", "텍스트 없음"
    except Exception as e:
        return f"[HWP 파싱 오류: {str(e)}]", "오류 발생"


def parse_image(file_path: str, ext: str) -> List[Dict[str, str]]:
    mime_map = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp"
    }
    mime_type = mime_map.get(ext, "image/jpeg")
    
    with open(file_path, "rb") as f:
        img_bytes = f.read()
        
    b64_str = base64.b64encode(img_bytes).decode("utf-8")
    return [{
        "mime_type": mime_type,
        "data": b64_str
    }]


def parse_plain_text(file_path: str) -> str:
    encodings = ["utf-8", "cp949", "euc-kr", "utf-16"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()
